import docx
import os

def generate_template():
    source_path = "トレーシングレポート_吸入_ひながた.docx"
    template_path = "template_kyunyu_report.docx"
    
    if not os.path.exists(source_path):
        print(f"Error: {source_path} not found.")
        return

    doc = docx.Document(source_path)

    replacements = {
        "相澤　一郎": "{{ doctor_name }}",
        "令和8年9月24日": "{{ report_date }}",
        "山田　太郎": "{{ patient_name }}",
        "平成29年3月14日": "{{ patient_dob }}",
        "あやめ薬局": "{{ pharmacy_name }}",
        "日立市南高野3丁目15番5号": "{{ pharmacy_address }}",
        "0294-33-5920": "{{ pharmacy_tel }}",
        "0294-33-5921": "{{ pharmacy_fax }}",
        "相澤　良太": "{{ pharmacist_name }}",
        "イナビル吸入粉末剤２０ｍｇ": "{{ target_drug }}",
    }

    def process_paragraph(p):
        full_text = "".join(r.text for r in p.runs)
        
        # 1. 患者からの同意（得た / 得ていない）の sdt コントロールを完全除去して置換
        if "得た" in full_text and "得ていない" in full_text:
            for sdt in p._p.xpath('.//w:sdt'):
                sdt.getparent().remove(sdt)
            for r in p.runs:
                r.text = ""
            if p.runs:
                p.runs[0].text = "　{{ consent_got }} 得た　　{{ consent_not_got }} 得ていない"
            return

        # 2. 拒否ですが の sdt コントロールを完全除去して置換
        if "拒否していますが" in full_text:
            for sdt in p._p.xpath('.//w:sdt'):
                sdt.getparent().remove(sdt)
            for r in p.runs:
                r.text = ""
            if p.runs:
                p.runs[0].text = "{{ consent_refused }} 患者は処方医への報告を拒否していますが、"
            return

        # 3. 報告日の「　日」重複対策
        if "報告日" in full_text:
            full_text = full_text.replace("　日", "").replace("日日", "日")
            full_text = full_text.replace("相澤　一郎", "{{ doctor_name }}")
            full_text = full_text.replace("令和8年9月24日", "{{ report_date }}")
            if p.runs:
                p.runs[0].text = full_text
                for r in p.runs[1:]:
                    r.text = ""
            return

        # 4. オーダー番号の置換
        if "オーダー番号" in full_text:
            if p.runs:
                p.runs[0].text = "オーダー番号：{{ order_no }}"
                for r in p.runs[1:]:
                    r.text = ""
            return

        # 5. 吸入手技の説明 が含まれる段落を {{ guidance_detail }} に置換
        if "吸入手技の説明" in full_text:
            if p.runs:
                p.runs[0].text = "{{ guidance_detail }}"
                for r in p.runs[1:]:
                    r.text = ""
            return

        # 6. 使用経験の確認 が含まれる段落は削除（空文字に）
        if "使用経験の確認" in full_text:
            if p.runs:
                for r in p.runs:
                    r.text = ""
            return

        # 単純な文字列置換
        for key, val in replacements.items():
            if key in full_text:
                full_text = full_text.replace(key, val)
                if p.runs:
                    p.runs[0].text = full_text
                    for r in p.runs[1:]:
                        r.text = ""

    for p in doc.paragraphs:
        process_paragraph(p)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    process_paragraph(p)

    doc.save(template_path)
    print(f"Fixed template created without sdt displacement: {template_path}")

if __name__ == "__main__":
    generate_template()
